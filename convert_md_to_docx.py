import docx

doc = docx.Document()

with open("Reyes-IT15_PROJECT FINAL DOCS.md", "r", encoding="utf-8") as f:
    for line in f:
        line = line.strip()
        if not line:
            doc.add_paragraph() # Add empty line
            continue
            
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif "**" in line:
            p = doc.add_paragraph()
            # Simple bold parsing
            parts = line.split("**")
            for i, part in enumerate(parts):
                if i % 2 == 1:
                    run = p.add_run(part)
                    run.bold = True
                else:
                    p.add_run(part)
        else:
            doc.add_paragraph(line)

doc.save("Reyes-IT15_PROJECT FINAL DOCS - V8.docx")
print("Successfully generated Reyes-IT15_PROJECT FINAL DOCS - V8.docx")
