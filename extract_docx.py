import docx
import sys

def extract_text(filename):
    doc = docx.Document(filename)
    full_text = []
    
    # Process paragraphs and tables in order (simplified, just all paras then all tables)
    for p in doc.paragraphs:
        if p.text.strip():
            full_text.append(p.text)
            
    for table in doc.tables:
        for row in table.rows:
            row_data = []
            for cell in row.cells:
                row_data.append(cell.text.replace("\n", " ").strip())
            full_text.append(" | ".join(row_data))
            
    with open("extracted_docx.txt", "w", encoding="utf-8") as f:
        f.write("\n".join(full_text))

if __name__ == "__main__":
    extract_text(sys.argv[1])
